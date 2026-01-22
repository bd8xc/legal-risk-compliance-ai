# DOCX Loader for extracting text and tables from DOCX files
import docx

def load_docx(file_path):
	doc = docx.Document(file_path)
	text_blocks = []
	tables_out = []

	# Extract paragraphs as text
	for para in doc.paragraphs:
		if para.text.strip():
			text_blocks.append(para.text.strip())

	# Extract tables
	for table in doc.tables:
		clean = []
		for row in table.rows:
			clean.append([cell.text.strip() for cell in row.cells])
		tables_out.append(clean)

	return {
		"text": "\n\n".join(text_blocks),
		"tables": tables_out
	}
