import pytest

pytestmark = pytest.mark.requires_docx

import io
from docx import Document
from ingestion.docx_loader import load_docx


def _make_docx(paragraphs=None, tables=None):
    doc = Document()
    for p in (paragraphs or []):
        doc.add_paragraph(p)
    for table_data in (tables or []):
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for i, row in enumerate(table_data):
            for j, cell in enumerate(row):
                table.cell(i, j).text = cell
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_load_docx_paragraphs():
    buf = _make_docx(paragraphs=["Hello world.", "Second paragraph."])
    result = load_docx(buf)
    assert "Hello world." in result["text"]
    assert "Second paragraph." in result["text"]


def test_load_docx_tables():
    table_data = [["Name", "Type"], ["Foo", "Bar"]]
    buf = _make_docx(tables=[table_data])
    result = load_docx(buf)
    assert len(result["tables"]) == 1
    assert result["tables"][0][0] == ["Name", "Type"]


def test_load_docx_empty():
    buf = _make_docx()
    result = load_docx(buf)
    assert result["text"] == ""
    assert result["tables"] == []


def test_load_docx_mixed():
    table_data = [["Col1", "Col2"]]
    buf = _make_docx(paragraphs=["Paragraph text."], tables=[table_data])
    result = load_docx(buf)
    assert "Paragraph text." in result["text"]
    assert len(result["tables"]) == 1
